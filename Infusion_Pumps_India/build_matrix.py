"""Builds Infusion_Pump_Classification_India.xlsx (B. Braun & Fresenius Kabi, India market)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Consulting/Infusion_Pumps_India/Infusion_Pump_Classification_India.xlsx"

F = "Arial"
NAVY = PatternFill("solid", fgColor="0B1F6B")
MID = PatternFill("solid", fgColor="3D5DA8")
BLUE = PatternFill("solid", fgColor="DEEBF7")      # listed on official India website
PEACH = PatternFill("solid", fgColor="FCE4D6")     # third-party (distributor) only
GREY = PatternFill("solid", fgColor="E0E0E0")      # not disclosed on Indian websites
RED = PatternFill("solid", fgColor="F8CBCB")       # not verified in India
INPUT = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
SRC_FILL = {"Official India website": BLUE, "Third-party (India distributor)": PEACH, "Not offered in India": GREY,
            "Not verified in India": RED}

# ---------------------------------------------------------------- product evidence
# (company, product, pump type, drug library evidence, connectivity evidence, classification,
#  India listing source, primary India URL, supporting URL(s), sample said, change vs sample, notes)
BB_CAT = "https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000529/automated-infusion-systems"
P = [
 ("B. Braun", "Perfusor compact", "Syringe",
  "No - not mentioned. Page lists rate range, volume pre-selection, syringe recognition only.",
  "None stated.", "Standard", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00000461",
  "Family page (fm Generation): https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000577/fm-generation",
  "Not shown (Standard column blank)", "ADD to Syringe-Standard", "Older fm-generation pump."),
 ("B. Braun", "Perfusor compact S", "Syringe",
  "No - not mentioned. Features: syringe recognition, bolus reduction, Data Lock (key lock, not a drug library).",
  "None stated.", "Standard", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00000145",
  "https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000577/fm-generation",
  "Not shown (Standard column blank)", "ADD to Syringe-Standard", "Transportable / home-care syringe pump."),
 ("B. Braun", "Perfusor compactplus", "Syringe",
  "Yes - 'integrated DoseGuard(TM) drug library' (Overview & Texts tab).",
  "Optional - pump links to hospital network only via separately sold 'Data module compactplus' (also listed on India catalogue: 'central interface to seamlessly link the compactplus infusion pumps to the hospital's network ... EMR'). Brochure p.5: module 'WiFi-connectivity included'.",
  "Smart**", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00010216",
  "Data module (India): https://catalogs.bbraun.co.in/en-IN/p/PRID00011150 | Brochure p.2/p.5/p.6: https://ecatalog.bbraun.com/eDoc/QlBSMDAwMDAwMDAwMDAwMDAwMTAwMDE4NTA5OTAwMDAw?openInline=true",
  "Smart", "RECLASSIFY Smart -> Smart** (connectivity optional via add-on module)", ""),
 ("B. Braun", "Perfusor Space", "Syringe",
  "Yes - 'Drug Library with capacity to up to 1200 drug names ... 30 different categories' (Overview & Texts > Advantages).",
  "Optional - 'B. Braun Space can also be integrated into the data communications network of every advanced hospital operation.'",
  "Smart**", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00001226",
  "Family page (Space System): https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000575/space-system",
  "Smart", "RECLASSIFY Smart -> Smart** (network integration optional)", "Click 'Read more' / 'Overview & Texts' tab to see drug library text."),
 ("B. Braun", "Spaceplus Perfusor", "Syringe",
  "Yes (global) - Spaceplus tech sheet p.2: '10,000 drugs including all parameters in drug library'.",
  "Integrated WiFi (global) - tech sheet p.1 battery time 'with WiFi activated'.",
  "Excluded - not offered in India", "Not offered in India",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00011858 (returns 'No results')",
  "India Spaceplus family lists only Spaceplus Infusomat: https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000574/spaceplus-system | Global: https://catalogs.bbraun.com/en-01/p/PRID00011858/spaceplus-perfusor",
  "Connected - 'Perfusor Space Plus' (official India website)", "REMOVED - not on B. Braun India catalogue; not found on Indian distributor sites (confirmed by project team, manual check)",
  "The shared Spaceplus technical-data PDF attached to the India Spaceplus Infusomat page covers both pumps - weak signal only, not a product listing."),
 ("B. Braun", "Infusomat fmS", "Volumetric",
  "No - not mentioned (dose-rate calculation only).", "None stated.", "Standard", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00000618",
  "https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000577/fm-generation",
  "Not shown (Standard column blank)", "ADD to Volumetric-Standard", ""),
 ("B. Braun", "Infusomat P", "Volumetric",
  "No - not mentioned (dose-rate mode, Data Lock only).", "None stated.", "Standard", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00000689",
  "https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000577/fm-generation",
  "Not shown (Standard column blank)", "ADD to Volumetric-Standard", ""),
 ("B. Braun", "Infusomat compactplus", "Volumetric",
  "Yes - 'integrated DoseGuard(TM) drug library'.",
  "Optional - via separately sold Data module compactplus (see Perfusor compactplus).",
  "Smart**", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00011042",
  "Data module (India): https://catalogs.bbraun.co.in/en-IN/p/PRID00011150 | Brochure p.5: https://ecatalog.bbraun.com/eDoc/QlBSMDAwMDAwMDAwMDAwMDAwMTAwMDE4NTA5OTAwMDAw?openInline=true",
  "Smart", "RECLASSIFY Smart -> Smart**", ""),
 ("B. Braun", "Infusomat compactplus P", "Volumetric",
  "Yes - 'possibility of colour coded drug libraries'; leaflet p.2 'Drug Library 3,000 drugs ... up to 30 drug categories'.",
  "Optional - via Data module compactplus.",
  "Smart**", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00012086",
  "Leaflet p.2: https://ecatalog.bbraun.com/eDoc/QlBSMDAwMDAwMDAwMDAwMDAwMTAwMDIyNDAwOTAwMDAw?openInline=true",
  "Not shown", "ADD to Volumetric-Smart**", "Variant for non-dedicated (standard gravity) IV sets."),
 ("B. Braun", "Infusomat Space P", "Volumetric",
  "Yes - 'Drug Library: Up to 720 drug names ... 15 categories'.",
  "Optional - 'B. Braun Space can also be integrated into the data communications network'.",
  "Smart**", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00001230",
  "Space System family (no plain 'Infusomat Space' listed): https://catalogs.bbraun.co.in/en-IN/c/PRODUCTS0000000575/space-system",
  "Smart - 'Infusomat Space'", "RENAME to Infusomat Space P (India-listed variant) + RECLASSIFY Smart -> Smart**", ""),
 ("B. Braun", "Spaceplus Infusomat", "Volumetric",
  "Yes - 'Drug Library Manager establish customized drug data bases ... dosage limits'.",
  "Integrated - 'stand-alone device with integrated WiFi'; 'Data communication towards EMR / PDMS systems via HL7 interface, Ethernet and WiFi'.",
  "Connected", "Official India website",
  "https://catalogs.bbraun.co.in/en-IN/p/PRID00011860",
  "Spaceplus brochure p.2: https://ecatalog.bbraun.com/eDoc/QlBSMDAwMDAwMDAwMDAwMDAwMTAwMDI2NDQyMTAwMDAw?openInline=true",
  "Connected - 'Infusomat Space Place'", "RENAME (typo) to Spaceplus Infusomat - classification confirmed", ""),
 ("Fresenius Kabi", "Infusia SP7s (SP7sED3)", "Syringe",
  "Yes - 'Syringe Infusion Pump with drug library - 17 therapy categories ... 1030 different drugs'.",
  "None stated on India page. (Third-party listings mention an RS232 port - service/data port, not treated as connectivity.)",
  "Smart", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/infusia-range/infusia-sp7sed3",
  "Distributor: https://arraymed.co.in/product/fresenius-kabi-sp7sed3/",
  "Smart", "CONFIRMED", "India ED3 version: 1,030 drugs/17 categories (MENA page shows older 201 drugs/14 categories)."),
 ("Fresenius Kabi", "Agilia SP", "Syringe",
  "India page silent. Distributor (Arraymed): 'Up to 19 customizable drug libraries' (listed as Agilia SP PCA).",
  "None stated on India page.",
  "Smart", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp",
  "Drug-library evidence (3rd party): https://arraymed.co.in/product/fresenius-kabi-agilia-sp/",
  "Not shown", "ADD to Syringe-Smart (drug-library evidence is third-party)", "Arraymed describes the PCA variant; FK India page does not specify variant."),
 ("Fresenius Kabi", "Agilia SP MC", "Syringe",
  "Yes - FK data sheet: 'Up to 19 embedded drug libraries' (Vigilant Drug'Lib); Arraymed: 'Up to 19 Drug Libraries'.",
  "Optional - India page: 'Agilia SP MC WiFi is the syringe pump in the Agilia Connect range that communicates via Wifi'; data sheet: 'Wireless LAN (For Agilia SP MC WiFi only)'.",
  "Smart**", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp-mc",
  "FK data sheet: https://www.fresenius-kabi.com/content/dam/fresenius-kabi/gb/products/product-documents/medtech/agilia-connect-sp/IFT265%20Agilia%20SP%20Connect%20Datasheet.pdf.coredownload.inline.pdf | https://arraymed.co.in/product/fresenius-kabi-agilia-spmc/",
  "Not shown", "ADD to Syringe-Smart**", ""),
 ("Fresenius Kabi", "Agilia SP MC WiFi", "Syringe",
  "Yes - same as Agilia SP MC (19 drug libraries).",
  "Integrated WiFi - 'communicates via Wifi with our Software Suite' (India page).",
  "Connected", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp-mc",
  "FK data sheet p.1-2 (link above)",
  "Connected", "CONFIRMED", "Shares the Agilia SP MC page on the India site."),
 ("Fresenius Kabi", "Agilia SP TIVA", "Syringe",
  "Yes - Arraymed: 'Up to 19 configurable drug libraries'.",
  "Optional - India page: 'Agilia Connect ... cyber secure communication'; Arraymed battery spec '13 Hours (Standard) / 9 Hours (WiFi)' = WiFi is a variant.",
  "Smart**", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp-tiva",
  "https://arraymed.co.in/product/fresenius-kabi-agilia-tiva-pump/",
  "Not shown", "ADD to Syringe-Smart**", "Anaesthesia (TCI) niche pump."),
 ("Fresenius Kabi", "Infusia VP7s (VP7sED3)", "Volumetric",
  "Yes - 'Volumetric Infusion Pump with drug library - 17 therapy categories ... 1030 different drugs'.",
  "None stated.", "Smart", "Official India website",
  "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/infusia-range/infusia-vp7sed3",
  "Distributor: https://arraymed.co.in/product/fresenius-kabi-vp7sed3/",
  "Smart", "CONFIRMED", ""),
 ("Fresenius Kabi", "Agilia VP", "Volumetric",
  "Yes - Arraymed: 'Soft & hard dose limit support' (drug-library limits).",
  "None stated.", "Smart", "Third-party (India distributor)",
  "https://arraymed.co.in/product/fresenius-kabi-agilia-vp/",
  "FK India Agilia range lists only SP, SP MC, SP TIVA: https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range",
  "Not shown", "ADD to Volumetric-Smart (peach - distributor only)", "Arraymed 'Download Brochure' button links to an unrelated Nihon Kohden defibrillator PDF - do not cite it."),
 ("Fresenius Kabi", "Agilia VP MC", "Volumetric",
  "Yes - FK data sheet: 'Up to 19 embedded drug libraries'; 'Drug library to be created with Agilia Vigilant Drug'Lib'.",
  "Optional - WiFi only on the separate Agilia VP MC WiFi variant (data sheet: 'Wireless LAN (For Agilia VP MC WiFi only)'). The WiFi variant was not found listed in India.",
  "Smart**", "Third-party (India distributor)",
  "https://dir.indiamart.com/search.mp?ss=agilia+vp+mc+wifi (listing: 'Fresenius Kabi Agilia VP MC'; product-page URL to be added)",
  "FK data sheet p.1-2: https://www.fresenius-kabi.com/content/dam/fresenius-kabi/gb/products/product-documents/medtech/agilia-connect-vp/IFT264%20Agilia%20VPMC%20Connect%20Data%20Sheet.pdf.coredownload.inline.pdf" ,
  "Connected - 'Agilia VP MC WiFi' (Indian distributors' websites)",
  "RECLASSIFY: India listing is base Agilia VP MC -> Volumetric-Smart**; WiFi variant not listed in India -> FK Volumetric-Connected now blank",
  "IndiaMart listing verified manually by project team (Oct-2026 screenshot); no price shown; seller not captured. Search for 'agilia vp mc wifi' returned no WiFi-variant listing."),
]

wb = Workbook()

def hdr(ws, row, values, fill=NAVY):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = fill; cell.alignment = CENTER; cell.border = BOX

def legend(ws, row, col):
    items = [(BLUE, "Listed on official India website"), (PEACH, "Listed on Indian distributor (third-party) website only"),
             (GREY, "Not disclosed on Indian websites"), (RED, "Pending confirmation of an Indian listing (none open)")]
    for i, (fill, txt) in enumerate(items):
        ws.cell(row=row + i, column=col).fill = fill
        ws.cell(row=row + i, column=col).border = BOX
        ws.cell(row=row + i, column=col + 1, value=txt).font = Font(name=F, size=9)

# ---------------------------------------------------------------- Sheet 1: slide matrix
ws = wb.active; ws.title = "Slide Matrix"
ws["A1"] = "Task 2.2 | Infusion Pumps Market Assessment in India | Key Players - corrected classification (B. Braun, Fresenius Kabi)"
ws["A1"].font = Font(name=F, bold=True, size=13, color="0B1F6B")
ws["A2"] = "Criteria: Standard = no drug library, no connectivity | Smart = drug library, no connectivity | Smart** = drug library, connectivity optional (add-on module / WiFi variant) | Connected = drug library + integrated connectivity"
ws["A2"].font = Font(name=F, italic=True, size=9)
legend(ws, 3, 11)
hdr(ws, 4, ["Company", "", "", "Syringe Pump", "", "", "Volumetric Pump", "", "", "Changes vs sample"])
ws.merge_cells("A4:C4"); ws.merge_cells("D4:F4"); ws.merge_cells("G4:I4")
hdr(ws, 5, ["Name", "HQ", "Revenue", "Standard", "Smart", "Connected", "Standard", "Smart", "Connected", ""], fill=MID)
ws.merge_cells("J4:J5")

COLS = {("Syringe", "Standard"): 4, ("Syringe", "Smart"): 5, ("Syringe", "Smart**"): 5, ("Syringe", "Connected"): 6,
        ("Volumetric", "Standard"): 7, ("Volumetric", "Smart"): 8, ("Volumetric", "Smart**"): 8, ("Volumetric", "Connected"): 9}
companies = [
 ("B. Braun", "Germany", "EUR 9.40B FY25 group sales (~USD 10.53B*)\nof which Hospital Care div. (incl. infusion therapy): EUR 5.09B\nSource: B. Braun Annual Report 2025",
  "- Standard columns were blank: add Perfusor compact, compact S, Infusomat fmS, Infusomat P\n- compactplus & Space pumps -> Smart** (connectivity is optional)\n- 'Perfusor Space Plus' removed: not offered in India (catalogue + distributors)\n- 'Infusomat Space' -> India lists Infusomat Space P\n- 'Infusomat Space Place' typo -> Spaceplus Infusomat\n- Add Infusomat compactplus P"),
 ("Fresenius Kabi", "Germany", "EUR 8.61B FY25 Kabi revenue (~USD 9.65B*)\nof which MedTech (incl. infusion pumps): EUR 1.61B\nSource: Fresenius Annual Report 2025",
  "- Add Agilia SP (Smart), Agilia SP MC & SP TIVA (Smart**)\n- Agilia VP (distributor only) -> Volumetric Smart, peach\n- 'Agilia VP MC WiFi' -> India listing (IndiaMart) is base Agilia VP MC -> Volumetric Smart** (peach); Volumetric Connected now blank\n- Infusia SP7s/VP7s & Agilia SP MC WiFi confirmed"),
]
r = 6
for name, hq, rev, changes in companies:
    buckets = {}
    for p in P:
        if p[0] != name or p[6] == "Not offered in India": continue
        cls = p[5].replace(" (global only)", "")
        col = COLS[(p[2], cls)]
        label = p[1] + ("**" if cls == "Smart**" else "") + (" (pending India confirmation)" if p[6] == "Not verified in India" else "")
        buckets.setdefault(col, []).append((label, SRC_FILL[p[6]]))
    n = max(len(v) for v in buckets.values())
    for c in range(1, 11):
        for i in range(n):
            ws.cell(row=r + i, column=c).border = BOX
    for c, v in ((1, name), (2, hq), (3, rev), (10, changes)):
        ws.cell(row=r, column=c, value=v).font = Font(name=F, size=9, bold=(c == 1))
        ws.cell(row=r, column=c).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=c, end_row=r + n - 1, end_column=c)
    for col in range(4, 10):
        items = buckets.get(col, [])
        for i in range(n):
            cell = ws.cell(row=r + i, column=col)
            if i < len(items):
                cell.value, cell.fill = items[i]
            else:
                cell.fill = GREY if not items else PatternFill()
            cell.font = Font(name=F, size=9); cell.alignment = CENTER
        if not items:
            ws.merge_cells(start_row=r, start_column=col, end_row=r + n - 1, end_column=col)
            ws.cell(row=r, column=col).fill = GREY
    for i in range(n): ws.row_dimensions[r + i].height = 30
    r += n
ws.cell(row=r + 1, column=1, value="** Smart** = drug library + connectivity optional (separately sold data module / WiFi variant).  "
        "* USD converted at ECB EUR/USD reference rate 1.1204 (5-Oct-2026), latest available; sample slide had used ~1.148.").font = Font(name=F, size=8, italic=True)
ws.cell(row=r + 2, column=1, value="BD and Baxter excluded from this pass per instruction. Product-level evidence and URLs: see 'Product Evidence' tab.").font = Font(name=F, size=8, italic=True)
for c, w in zip("ABCDEFGHIJKL", [14, 9, 30, 18, 22, 22, 18, 26, 24, 55, 4, 48]):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "A6"

# ---------------------------------------------------------------- Sheet 2: product evidence
ws2 = wb.create_sheet("Product Evidence")
heads = ["Company", "Product", "Pump type", "Drug library - evidence (quoted)", "Connectivity - evidence (quoted)",
         "Classification", "India listing source", "Primary URL (India, client-viewable)", "Supporting URL(s) / brochure page",
         "Sample slide said", "Change vs sample", "Notes", "Evidence screenshot (evidence/ folder)"]
hdr(ws2, 1, heads)
shots = {"Perfusor compact": "BBraun_IN_Perfusor_compact.jpg", "Perfusor compact S": "BBraun_IN_Perfusor_compact_S.jpg",
         "Perfusor compactplus": "BBraun_IN_Perfusor_compactplus.jpg; BBraun_IN_Data_module_compactplus.jpg",
         "Perfusor Space": "BBraun_IN_Perfusor_Space.jpg", "Spaceplus Perfusor": "BBraun_IN_Spaceplus_Perfusor_NO_RESULTS.jpg",
         "Infusomat fmS": "BBraun_IN_Infusomat_fmS.jpg", "Infusomat P": "BBraun_IN_Infusomat_P.jpg",
         "Infusomat compactplus": "BBraun_IN_Infusomat_compactplus.jpg", "Infusomat compactplus P": "BBraun_IN_Infusomat_compactplus_P.jpg",
         "Infusomat Space P": "BBraun_IN_Infusomat_Space_P.jpg", "Spaceplus Infusomat": "BBraun_IN_Spaceplus_Infusomat.jpg",
         "Infusia SP7s (SP7sED3)": "FK_IN_Infusia_SP7s.jpg", "Agilia SP": "FK_IN_Agilia_SP.jpg; Arraymed_Agilia_SP.jpg",
         "Agilia SP MC": "FK_IN_Agilia_SP_MC.jpg; Arraymed_Agilia_SP_MC.jpg", "Agilia SP MC WiFi": "FK_IN_Agilia_SP_MC.jpg",
         "Agilia SP TIVA": "FK_IN_Agilia_SP_TIVA.jpg; Arraymed_Agilia_SP_TIVA.jpg", "Infusia VP7s (VP7sED3)": "FK_IN_Infusia_VP7s.jpg",
         "Agilia VP": "Arraymed_Agilia_VP.jpg; FK_IN_Agilia_range.jpg", "Agilia VP MC": "IndiaMart_search_Agilia_VP_MC_team_check.jpg"}
for i, p in enumerate(P, 2):
    row = list(p) + [shots.get(p[1], "")]
    for c, v in enumerate(row, 1):
        cell = ws2.cell(row=i, column=c, value=v)
        cell.font = Font(name=F, size=9, bold=(c == 6)); cell.alignment = WRAP; cell.border = BOX
    fill = SRC_FILL[p[6]]
    for c in (2, 7, 8): ws2.cell(row=i, column=c).fill = fill
    if p[7].startswith("http"): ws2.cell(row=i, column=8).hyperlink = p[7]; ws2.cell(row=i, column=8).font = Font(name=F, size=9, color="0563C1", underline="single")
for c, w in enumerate([13, 22, 11, 42, 48, 13, 18, 44, 55, 22, 32, 40, 30], 1):
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.freeze_panes = "C2"; ws2.auto_filter.ref = f"A1:M{len(P)+1}"

# ---------------------------------------------------------------- Sheet 3: revenue
ws3 = wb.create_sheet("Revenue")
ws3["A1"] = "Revenue cross-check (FY2025)"; ws3["A1"].font = Font(name=F, bold=True, size=12)
ws3["A2"] = "EUR/USD rate (input)"; ws3["B2"] = 1.1204; ws3["B2"].font = Font(name=F, color="0000FF"); ws3["B2"].fill = INPUT
ws3["C2"] = "ECB euro reference rate, latest available (5-Oct-2026): https://data.ecb.europa.eu/data/datasets/EXR/EXR.D.USD.EUR.SP00.A  |  For reference: 2025 annual average 1.1300; the sample slide implied ~1.148."
for c in ("A2", "C2"): ws3[c].font = Font(name=F, size=9, italic=(c == "C2"))
hdr(ws3, 4, ["Company", "Metric", "Period", "Value (EUR mn)", "Value (USD bn) - formula", "Value (INR, as published)",
             "Sample slide", "Check", "Source URL", "Page / location", "Comment"])
rev = [
 ("B. Braun", "Group sales", "FY2025 (Dec-25)", 9396, None, "~$10.79B (FY25, Group Sales)", "EUR figure matches; sample USD used ~1.148 -> now USD 10.53B at latest ECB rate",
  "https://www.bbraun.com/en/about-us/company/facts-and-figures/annual-report.html", "Key figures table; Annual Report PDF p.8",
  "PDF: https://www.bbraun.com/content/dam/b-braun/master/website-6/en/04_about-us/0401_company/facts-and-figures/2025_B_Braun_Annual_Report.pdf"),
 ("B. Braun", "Hospital Care division sales (incl. infusion therapy)", "FY2025", 5090.8, None, "-", "Closest infusion proxy",
  "https://www.bbraun.com/content/dam/b-braun/master/website-6/en/04_about-us/0401_company/facts-and-figures/2025_B_Braun_Annual_Report.pdf",
  "PDF p.43 (section 4.9); p.8 'Hospital Care 5,091 ... infusion therapy, nutrition therapy and pain therapy'",
  "Pump-only revenue not disclosed. AR p.41 notes protectionist procurement in China and India affected foreign medical device sales."),
 ("B. Braun", "B. Braun Medical (India) Pvt Ltd - revenue", "FY2025 (Mar-25)", None, None, "INR 500-750 Cr", "UNOFFICIAL - not used on slide",
  "https://www.tofler.in/b-braun-medical-india-private-limited/company/U33112MH1984PTC214514", "Key metrics (range; exact figure paywalled)",
  "Tofler: revenue growth 13.34%. Tracxn range INR 500-1,000 Cr: https://tracxn.com/d/legal-entities/india/b.braun-medical-india-private-limited/__bN1KxXSC4dkyGIflutZRBCshiSEXqK9hrOWP4hHQsdA"),
 ("Fresenius Kabi", "Fresenius Kabi segment revenue", "FY2025 (Dec-25)", 8612, None, "~$9.89B (FY25)", "EUR figure matches; sample USD used ~1.148 -> now USD 9.65B at latest ECB rate",
  "https://report.fresenius.com/2025/annual-report/financial-statements/segment-reporting.html", "Segment table; Annual Report PDF p.5 and p.117",
  "PDF: https://www.fresenius.com/sites/default/files/2026-03/fresenius_annual_report_2025_0.pdf"),
 ("Fresenius Kabi", "MedTech business revenue (incl. infusion pumps)", "FY2025", 1610, None, "-", "Closest infusion proxy",
  "https://www.fresenius.com/sites/default/files/2026-03/fresenius_annual_report_2025_0.pdf", "PDF p.117 (printed p.116): 'Revenue in the MedTech business increased ... to EUR 1,610 million'",
  "Also includes transfusion, disposables and Ivenix (US). Pump-only revenue not disclosed."),
 ("Fresenius Kabi", "Fresenius Kabi India Pvt Ltd - revenue", "FY2025 (Mar-25)", None, None, "INR 750-1,000 Cr", "UNOFFICIAL - not used on slide",
  "https://www.tofler.in/fresenius-kabi-india-private-limited/company/U24231PN1995PTC014017", "Key metrics (range; exact figure paywalled)",
  "Tofler revenue growth 9.83%. Entity covers pumps, disposables, nutrition, oncology etc. (Fresenius Kabi Oncology Ltd is a separate entity.)"),
]
for i, rw in enumerate(rev, 5):
    vals = list(rw[:4]) + [f"=IF(D{i}=\"\",\"n/a\",D{i}*$B$2/1000)"] + list(rw[5:])
    for c, v in enumerate(vals, 1):
        cell = ws3.cell(row=i, column=c, value=v); cell.font = Font(name=F, size=9); cell.alignment = WRAP; cell.border = BOX
    ws3.cell(row=i, column=4).font = Font(name=F, size=9, color="0000FF"); ws3.cell(row=i, column=4).number_format = "#,##0.0"
    ws3.cell(row=i, column=5).number_format = "0.00"
    ws3.cell(row=i, column=9).hyperlink = rw[8]
for c, w in enumerate([14, 30, 15, 13, 14, 18, 22, 22, 50, 38, 55], 1):
    ws3.column_dimensions[get_column_letter(c)].width = w

# ---------------------------------------------------------------- Sheet 4: clean URL map
ws4 = wb.create_sheet("Clean URL Map")
hdr(ws4, 1, ["Company", "Product / purpose", "URL (client-facing)", "Source type", "Status", "Reason / note"])
urls = [
 ("B. Braun", "India catalogue - automated infusion systems (master list)", BB_CAT, "Official India website", "KEEP", "Shows the 4 families listed in India"),
 ("B. Braun", "Perfusor compact", "https://catalogs.bbraun.co.in/en-IN/p/PRID00000461", "Official India website", "ADDED", "Standard"),
 ("B. Braun", "Perfusor compact S", "https://catalogs.bbraun.co.in/en-IN/p/PRID00000145", "Official India website", "ADDED", "Standard"),
 ("B. Braun", "Perfusor compactplus", "https://catalogs.bbraun.co.in/en-IN/p/PRID00010216", "Official India website", "REPLACES global link", "India equivalent of catalogs.bbraun.com/en-01/p/PRID00010216"),
 ("B. Braun", "Data module compactplus (connectivity add-on)", "https://catalogs.bbraun.co.in/en-IN/p/PRID00011150", "Official India website", "ADDED", "Proves connectivity is optional -> Smart**"),
 ("B. Braun", "compactplus brochure (p.5 data module, p.6 DoseGuard)", "https://ecatalog.bbraun.com/eDoc/QlBSMDAwMDAwMDAwMDAwMDAwMTAwMDE4NTA5OTAwMDAw?openInline=true", "Brochure linked from India catalogue", "ADDED", ""),
 ("B. Braun", "Perfusor Space", "https://catalogs.bbraun.co.in/en-IN/p/PRID00001226", "Official India website", "ADDED", "Drug library under 'Overview & Texts' tab"),
 ("B. Braun", "Infusomat fmS", "https://catalogs.bbraun.co.in/en-IN/p/PRID00000618", "Official India website", "ADDED", "Standard"),
 ("B. Braun", "Infusomat P", "https://catalogs.bbraun.co.in/en-IN/p/PRID00000689", "Official India website", "ADDED", "Standard"),
 ("B. Braun", "Infusomat compactplus", "https://catalogs.bbraun.co.in/en-IN/p/PRID00011042", "Official India website", "REPLACES global link", "India equivalent of catalogs.bbraun.com/en-01/p/PRID00011042"),
 ("B. Braun", "Infusomat compactplus P", "https://catalogs.bbraun.co.in/en-IN/p/PRID00012086", "Official India website", "ADDED", ""),
 ("B. Braun", "Infusomat Space P", "https://catalogs.bbraun.co.in/en-IN/p/PRID00001230", "Official India website", "ADDED", ""),
 ("B. Braun", "Spaceplus Infusomat", "https://catalogs.bbraun.co.in/en-IN/p/PRID00011860", "Official India website", "REPLACES global link", "India equivalent of catalogs.bbraun.com/en-01/p/PRID00011860"),
 ("B. Braun", "Annual Report 2025 (revenue)", "https://www.bbraun.com/en/about-us/company/facts-and-figures/annual-report.html", "Company annual report", "KEEP", "Group sales EUR 9,396mn"),
 ("B. Braun", "Annual Report 2025 PDF (p.8, p.43)", "https://www.bbraun.com/content/dam/b-braun/master/website-6/en/04_about-us/0401_company/facts-and-figures/2025_B_Braun_Annual_Report.pdf", "Company annual report", "ADDED", "Hospital Care division sales"),
 ("B. Braun", "India entity financials", "https://www.tofler.in/b-braun-medical-india-private-limited/company/U33112MH1984PTC214514", "Third-party data aggregator", "ADDED", "Revenue range only"),
 ("B. Braun", "compactplus system (global page)", "https://www.bbraun.com/en/products-and-solutions/therapies/infusion-therapy/automated-infusion-systems/compactplus-system.html", "Global website", "DROP", "Global page; India catalogue pages cover the same products"),
 ("B. Braun", "Spaceplus Perfusor (global catalogue)", "https://catalogs.bbraun.com/en-01/p/PRID00011858/spaceplus-perfusor", "Global website", "DROP", "Not listed on India catalogue (India URL returns 'No results')"),
 ("Fresenius Kabi", "India - Infusion Therapy (navigation)", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy", "Official India website", "KEEP", "Navigation only"),
 ("Fresenius Kabi", "India - Agilia range (lists SP, SP MC, SP TIVA only)", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range", "Official India website", "ADDED", "Shows no Agilia volumetric pump on India site"),
 ("Fresenius Kabi", "India - Infusia range", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/infusia-range", "Official India website", "KEEP", ""),
 ("Fresenius Kabi", "Infusia SP7s", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/infusia-range/infusia-sp7sed3", "Official India website", "KEEP", "Duplicate in source list removed"),
 ("Fresenius Kabi", "Infusia VP7s", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/infusia-range/infusia-vp7sed3", "Official India website", "KEEP", ""),
 ("Fresenius Kabi", "Agilia SP", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp", "Official India website", "ADDED", ""),
 ("Fresenius Kabi", "Agilia SP MC / SP MC WiFi", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp-mc", "Official India website", "KEEP", "Covers both SP MC and SP MC WiFi"),
 ("Fresenius Kabi", "Agilia SP TIVA", "https://www.fresenius-kabi.com/in/products/ins/infusion-therapy/agilia-range/agilia-sp-tiva", "Official India website", "ADDED", ""),
 ("Fresenius Kabi", "Agilia VP (distributor)", "https://arraymed.co.in/product/fresenius-kabi-agilia-vp/", "Third-party (India distributor)", "KEEP (peach)", "Ignore the page's brochure button - wrong PDF"),
 ("Fresenius Kabi", "Agilia SP (drug-library evidence)", "https://arraymed.co.in/product/fresenius-kabi-agilia-sp/", "Third-party (India distributor)", "ADDED (peach)", ""),
 ("Fresenius Kabi", "Agilia SP MC (drug-library evidence)", "https://arraymed.co.in/product/fresenius-kabi-agilia-spmc/", "Third-party (India distributor)", "ADDED (peach)", ""),
 ("Fresenius Kabi", "Agilia SP TIVA (WiFi variant evidence)", "https://arraymed.co.in/product/fresenius-kabi-agilia-tiva-pump/", "Third-party (India distributor)", "ADDED (peach)", ""),
 ("Fresenius Kabi", "Agilia SP MC / SP MC WiFi data sheet", "https://www.fresenius-kabi.com/content/dam/fresenius-kabi/gb/products/product-documents/medtech/agilia-connect-sp/IFT265%20Agilia%20SP%20Connect%20Datasheet.pdf.coredownload.inline.pdf", "Company brochure (global)", "ADDED (backup)", "'Wireless LAN (For Agilia SP MC WiFi only)'; 19 drug libraries"),
 ("Fresenius Kabi", "Segment reporting FY2025 (revenue)", "https://report.fresenius.com/2025/annual-report/financial-statements/segment-reporting.html", "Company annual report", "KEEP", "Kabi revenue EUR 8,612mn (duplicate in source list removed)"),
 ("Fresenius Kabi", "Annual Report 2025 PDF (p.5, p.117)", "https://www.fresenius.com/sites/default/files/2026-03/fresenius_annual_report_2025_0.pdf", "Company annual report", "KEEP", "p.117 for MedTech EUR 1,610mn"),
 ("Fresenius Kabi", "India entity financials", "https://www.tofler.in/fresenius-kabi-india-private-limited/company/U24231PN1995PTC014017", "Third-party data aggregator", "ADDED", "Revenue range only"),
 ("Fresenius Kabi", "GB product literature library", "https://www.fresenius-kabi.com/gb/healthcare-professional-area/medtech/product-literature-library", "Global website (GB HCP area)", "DROP", "GB HCP-only area; use the direct data-sheet PDF instead"),
 ("Fresenius Kabi", "Key2 Declaration of Conformity", "https://key2.fresenius-kabi.com/en/regulatory-document/declaration-conformity/mdr/declaration-conformity", "Regulatory document (EU)", "DROP", "EU regulatory document - proves variants exist (SP, SP MC, SP MC WiFi, SP TIVA WiFi) but not India availability"),
 ("Fresenius Kabi", "Agilia VP MC (IndiaMart search; team-verified listing)", "https://dir.indiamart.com/search.mp?ss=agilia+vp+mc+wifi", "Third-party (India distributor)", "ADDED (peach)", "Replace with the product-page URL of the 'Fresenius Kabi Agilia VP MC' listing"),
 ("Fresenius Kabi", "Agilia VP MC / VP MC WiFi data sheet", "https://www.fresenius-kabi.com/content/dam/fresenius-kabi/gb/products/product-documents/medtech/agilia-connect-vp/IFT264%20Agilia%20VPMC%20Connect%20Data%20Sheet.pdf.coredownload.inline.pdf", "Company brochure (global)", "ADDED (backup)", "'Wireless LAN (For Agilia VP MC WiFi only)'; 19 drug libraries"),
 ("BD / Baxter", "All BD and Baxter URLs", "-", "-", "OUT OF SCOPE", "Skipped per instruction"),
]
for i, u in enumerate(urls, 2):
    for c, v in enumerate(u, 1):
        cell = ws4.cell(row=i, column=c, value=v); cell.font = Font(name=F, size=9); cell.alignment = WRAP; cell.border = BOX
    if u[2].startswith("http"):
        ws4.cell(row=i, column=3).hyperlink = u[2]; ws4.cell(row=i, column=3).font = Font(name=F, size=9, color="0563C1", underline="single")
    if "Third-party (India" in u[3]: ws4.cell(row=i, column=4).fill = PEACH
    elif "Official India" in u[3]: ws4.cell(row=i, column=4).fill = BLUE
    if u[4] == "DROP": ws4.cell(row=i, column=5).fill = RED
for c, w in enumerate([14, 40, 70, 26, 18, 55], 1):
    ws4.column_dimensions[get_column_letter(c)].width = w
ws4.freeze_panes = "A2"; ws4.auto_filter.ref = f"A1:F{len(urls)+1}"

# ---------------------------------------------------------------- Sheet 5: method & caveats
ws5 = wb.create_sheet("Method & Caveats")
notes = [
 "Scope: India market, B. Braun and Fresenius Kabi; syringe and volumetric (large-volume) pumps. BD and Baxter skipped per instruction. Checked October 2026.",
 "Source priority: (1) company India website, including its catalogue tabs and linked brochures; (2) Indian distributor websites incl. IndiaMart, marked peach; (3) global company pages/data sheets, used only as backup for features and never as proof of India availability.",
 "Visual check: every India page was opened in a real browser and screenshotted (see evidence/ folder). B. Braun catalogue details sit under 'Read more' / 'Overview & Texts' - click there to see drug-library text.",
 "Classification rules applied: an add-on connectivity module (e.g., Data module compactplus) or a separate WiFi variant (e.g., Agilia SP MC vs SP MC WiFi) = 'connectivity optional' = Smart**. A basic RS232 / nurse-call port is NOT counted as connectivity.",
 "Where a product page is silent on the drug library (e.g., FK Agilia SP, SP TIVA), distributor text was used and the evidence column says so.",
 "LIMITATION: IndiaMart blocks automated access from this environment (HTTP 429). IndiaMart checks were done manually by the project team (screenshot in evidence/).",
 "LIMITATION: Revenue for the India entities is shown only as ranges (Tofler/Tracxn); exact MCA figures are paywalled. Neither group reports infusion-pump-only revenue; Hospital Care (B. Braun) and MedTech (Fresenius Kabi) are the closest disclosed proxies.",
 "Revenue in USD uses the latest ECB EUR/USD reference rate 1.1204 (5-Oct-2026) - an editable input on the Revenue tab. Slide shows official (annual report) figures only; India-entity ranges come from aggregators (Tofler/Tracxn) and are kept off the slide. Official India figures would need MCA filings (paid access).",
]
ws5["A1"] = "Method, assumptions and caveats"; ws5["A1"].font = Font(name=F, bold=True, size=12)
for i, n in enumerate(notes, 3):
    ws5.cell(row=i, column=1, value=f"{i-2}. {n}").font = Font(name=F, size=10)
    ws5.cell(row=i, column=1).alignment = WRAP
    ws5.row_dimensions[i].height = 42
ws5.column_dimensions["A"].width = 150

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
