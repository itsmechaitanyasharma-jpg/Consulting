"""Slide 2/8 - Terumo, Mindray, ICU Medical: India infusion pump classification workbook."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Consulting/Infusion_Pumps_India/Infusion_Pump_Classification_India_Slide2.xlsx"

F = "Arial"
NAVY = PatternFill("solid", fgColor="0B1F6B")
MID = PatternFill("solid", fgColor="3D5DA8")
BLUE = PatternFill("solid", fgColor="DEEBF7")      # listed on official India website
PEACH = PatternFill("solid", fgColor="FCE4D6")     # third-party (distributor) only
GREY = PatternFill("solid", fgColor="E0E0E0")      # not disclosed on Indian websites
RED = PatternFill("solid", fgColor="F8CBCB")       # not verified in India
INPUT = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
SRC_FILL = {"Official India website": BLUE, "Third-party (India distributor)": PEACH, "Not offered in India": GREY,
            "Not verified in India": RED}

# ---------------------------------------------------------------- sources
TER_PAGE = "https://terumoindia.com/medication-management/terufusiontm-advanced-infusion-systems"
TER_BRO = "https://terumoindia.com/sites/default/files/2025-11/Smart%20Infusion%20System.pdf"
TER_PR = "https://terumoindia.com/terufusion-connected-pumps"
MR = "https://www.mindray.com/in/products/infusion-system/"
MR1, MR3, MR5 = MR + "benefusion-1-series/", MR + "benefusion-3-series/", MR + "benefusion-5-series/"
MRE, MRU = MR + "benefusion-e-series/", MR + "benefusion-u-series/"
MR_SP3_DS = "https://5.imimg.com/data5/SELLER/Doc/2022/1/TX/CI/VD/17862303/mindray-benefusion-sp3-syringe-pump.pdf"
MR_U_BRO = "http://5.imimg.com/data5/SELLER/Doc/2025/12/571961803/KS/IN/JK/2785113/infusion-pump-uvp.pdf"
MR_USP_DS = "https://ecomed.co.za/wp-content/uploads/2025/03/BeneFusion-uSP_Data-Sheet_ENG_230120.pdf"
MR_ESP_DS = "https://ecomed.co.za/wp-content/uploads/2025/03/Syringe-driver-PCA-BeneFusion-eSP-PCA_Data-Sheet_ENG_20241203.pdf"
MR_VP5_OM = "http://mindray.sy/wp-content/uploads/2019/12/VP5-OperatorManual.pdf"
ICU_MF = "https://www.icumed.com/products/infusion-therapy/infusion-pumps-and-software/medfusion-syringe-pumps/medfusion-4000-wireless-syringe-infusion-pump/"
IM_MF = "https://www.indiamart.com/proddetail/mri-compatible-syringe-infusion-pump-medfusion-4000-21117047897.html"
IM_G1200 = "https://www.indiamart.com/proddetail/smiths-medical-graseby-1200-volumetric-infusion-pump-2853092683762.html"
MB_G1200 = "https://www.medikabazaar.com/products/smiths-medical-graseby-1200-infusion-pump-mbpgmdcispisff16979"
G1200_BRO = "http://www.hospimax.com/medical-devices/smithsmedical/Graseby%201200%20LVP%20Brochure%20(FINAL).pdf"
IM_G2000 = "https://www.indiamart.com/proddetail/smiths-medical-syringe-pump-graseby-2000-2852038109462.html"
IM_G2100 = "https://www.indiamart.com/proddetail/smith-syringe-infusion-pump-graseby-2100-21099225197.html"
HON_G2100 = "https://www.honmed.in/product-page/syringe-pump-graseby-2100"

# ---------------------------------------------------------------- product evidence
# (company, product, type, drug library evidence, connectivity evidence, classification, India source,
#  primary URL, supporting URLs, sample said, change vs sample, notes)
P = [
 ("Terumo", "TERUFUSION Syringe Pump Type SS3 TE-SS730", "Syringe",
  "None. Brochure spec table: 'Library mode (TE-SS830 only)'.",
  "None. Brochure p.2 groups TE-SS730 under 'Standard Pumps' - 'Standard pumps without IT functions'. RS-232C port on TE-SS732 variant only (not counted as connectivity).",
  "Standard", "Official India website", TER_BRO,
  "Brochure p.2 (Smart vs Standard panel) and p.4 (spec table). India product page: " + TER_PAGE,
  "Smart", "RECLASSIFY Smart -> Standard (no drug library, no IT function)",
  "Shown only in the brochure hosted on terumoindia.com (brochure dated Aug-2018). The India web page text names only the Smart pumps - confirm TE-SS730 is still actively sold in India."),
 ("Terumo", "TERUFUSION Syringe Pump Type SS3 TE-SS830", "Syringe",
  "Brochure p.4 'Smart Pumps [TE-LM830 / TE-SS830] only' - 'Drug library function minimizes human error'; spec: 'Library mode (TE-SS830 only)'. India page: 'integrated drug library'.",
  "Brochure spec table: 'External communication function (wireless LAN)' - TE-SS830 only; 'Links with Hospital Information Systems (HIS)'. India page: 'remote monitoring', Pump Monitoring System (PMS) software.",
  "Connected", "Official India website", TER_PAGE,
  "Brochure p.2 & p.4: " + TER_BRO + " | Launch release (Aug-2025): " + TER_PR,
  "Connected", "Confirmed", "Wireless LAN is built into the 830 models (the 730 models have none)."),
 ("Terumo", "TERUFUSION Infusion Pump Type LM3 TE-LM730", "Volumetric",
  "None. Brochure spec table: 'Library mode (TE-LM830 only)'.",
  "None. Brochure p.2 groups TE-LM730 under 'Standard Pumps' - 'without IT functions'. RS-232C on TE-LM732 variants only (not counted).",
  "Standard", "Official India website", TER_BRO,
  "Brochure p.2 and p.4. India product page: " + TER_PAGE,
  "Smart", "RECLASSIFY Smart -> Standard (no drug library, no IT function)",
  "Same caveat as TE-SS730: confirm it is still actively sold in India."),
 ("Terumo", "TERUFUSION Infusion Pump Type LM3 TE-LM830", "Volumetric",
  "Brochure: 'Library mode (TE-LM830 only)'; Drug Library Manager TE-SW800B 'TE-LM830 only'. India page: 'integrated drug library'.",
  "Brochure spec table: 'External communication function (wireless LAN)' - TE-LM830 only. India page: remote monitoring via PMS software.",
  "Connected", "Official India website", TER_PAGE,
  "Brochure p.2 & p.4: " + TER_BRO + " | Launch release: " + TER_PR,
  "Connected", "Confirmed", ""),
 ("Mindray", "BeneFusion 1 Series (SP1)", "Syringe",
  "None mentioned on India page (features: compact, fast start, battery, IP34, EN-1789).",
  "None mentioned.", "Standard", "Official India website", MR1, "", "Standard", "Confirmed",
  "India page shows both a syringe and a volumetric pump in the hero image."),
 ("Mindray", "BeneFusion 1 Series (VP1)", "Volumetric",
  "None mentioned on India page.", "None mentioned.", "Standard", "Official India website", MR1, "", "Standard", "Confirmed", ""),
 ("Mindray", "BeneFusion 3 Series (SP3)", "Syringe",
  "India page: none mentioned. Data sheet hosted on IndiaMart: 'Drug library up to 200 drugs; ON/OFF switchable' - a drug-name list; no dose limits (DERS) stated.",
  "None. India page: docking station only (power/cable management).",
  "Standard", "Official India website", MR3, "SP3 data sheet (IndiaMart-hosted): " + MR_SP3_DS,
  "Standard", "Kept Standard - DECISION NEEDED (see note)",
  "Judgement call: if a drug-NAME library without dose limits counts as a 'drug library', SP3/VP3 move to Smart. Kept Standard, as in the sample, because there are no dose limits (DERS)."),
 ("Mindray", "BeneFusion 3 Series (VP3)", "Volumetric",
  "India page: none mentioned. VP3 data sheet: 'Drug library up to 200 drugs; ON/OFF switchable' - no dose limits stated.",
  "None.", "Standard", "Official India website", MR3, "India page hero/product images show volumetric + syringe 3 Series pumps.",
  "Standard", "Kept Standard - DECISION NEEDED (see note)", "Same judgement call as SP3."),
 ("Mindray", "BeneFusion 5 Series (SP5)", "Syringe",
  "India page: 'Dose Error Reduction System (DERS) - The DERS Drug Library ... alert ... when the soft/hard limit is reached'.",
  "India page: connectivity via 'DS5 Docking Station' and 'BeneFusion CS5 Central Infusion Supervision System'. VP5 operator manual s.7.4: 'Wireless Networking (Optional)'; 'The wireless module ... [is] optional'.",
  "Smart**", "Official India website", MR5, "VP5 operator manual s.7.4: " + MR_VP5_OM,
  "Smart", "RECLASSIFY Smart -> Smart** (connectivity via optional wireless module / DS5 dock)",
  "India page also lists 'SP5 TCI/DTCI' - a software mode of SP5 (not counted separately, same as Perfusor Space TCI)."),
 ("Mindray", "BeneFusion 5 Series (VP5)", "Volumetric",
  "India page: DERS Drug Library (5 Series).",
  "VP5 operator manual s.7.4: 'Wireless Networking (Optional)'. India page: DS5 dock / CS5 central supervision.",
  "Smart**", "Official India website", MR5, "VP5 operator manual s.7.4: " + MR_VP5_OM,
  "Smart", "RECLASSIFY Smart -> Smart**", "VP5 is not named in the India page text; the 5 Series product images show the pump stack."),
 ("Mindray", "BeneFusion e Series (eSP)", "Syringe",
  "India page: 'SafeDose DERS helps prevent dosing error with hard or soft limits restriction'. Data sheet: 'Drug library up to 5000 drugs'.",
  "India page: 'Integrated Central Monitoring - BeneVision CMS offers one-stop monitoring of all patients' vital sign and infusion treatment details'. eSP data sheet: 'Communication: Wired/wireless' (not marked optional).",
  "Connected", "Official India website", MRE, "eSP (PCA) data sheet, Connectivity section: " + MR_ESP_DS,
  "Smart", "RECLASSIFY Smart -> Connected (drug library + wired/wireless central monitoring)",
  "Data sheet is from a South African distributor (global Mindray document); the India page shows connectivity on its own."),
 ("Mindray", "BeneFusion e Series (eVP)", "Volumetric",
  "India page: SafeDose DERS (hard/soft limits).",
  "India page: Integrated Central Monitoring (BeneVision CMS). e Series data sheets: 'Communication: Wired/wireless'.",
  "Connected", "Official India website", MRE, "e Series brochure / data sheets (global)", "Smart",
  "RECLASSIFY Smart -> Connected", ""),
 ("Mindray", "BeneFusion u Series (uSP / uDSP)", "Syringe",
  "u Series brochure (IndiaMart-hosted) p.2: 'Departmental drug library - Customized drug library for varied care units'. uSP data sheet: 'Drug library up to 5000 drugs'; DERS optional.",
  "India page: 'Comprehensive integration, better connectivity ... Access the HIS/CIS/EMR system through the HL7 protocol'. uSP data sheet: 'Communication: Wired/wireless'.",
  "Connected", "Official India website", MRU, "u Series brochure (IndiaMart-hosted): " + MR_U_BRO + " | uSP data sheet: " + MR_USP_DS,
  "Connected", "Confirmed", "The India page itself does not mention the drug library; that evidence comes from the brochure."),
 ("Mindray", "BeneFusion u Series (uVP)", "Volumetric",
  "u Series brochure (IndiaMart-hosted) p.2: 'Departmental drug library'.",
  "India page: HL7 connectivity to HIS/CIS/EMR; centralised monitoring.",
  "Connected", "Official India website", MRU, "u Series brochure (IndiaMart-hosted): " + MR_U_BRO, "Connected", "Confirmed", ""),
 ("ICU Medical", "Graseby 2000", "Syringe",
  "None. Distributor text: rate / VTBI programming, '4-step programming sequence'.",
  "None.", "Standard", "Third-party (India distributor)", IM_G2000,
  "Honmed (Graseby 2000 range text): " + HON_G2100, "Standard (peach)", "Confirmed",
  "IndiaMart listing verified by project team (Oct-2026): seller Nature's Global Service, New Delhi, Rs 30,000; image shows a Graseby 2000."),
 ("ICU Medical", "Graseby 2100", "Syringe",
  "None. Honmed: '2100 ... also features Bodyweight Programming' (dose calculation, not a drug library).",
  "None.", "Standard", "Third-party (India distributor)", HON_G2100, "IndiaMart: " + IM_G2100, "Standard (peach)", "Confirmed",
  "IndiaMart listing verified by project team: seller Apex Medical India (Mark.Inc), Lucknow, Rs 30,000; pump screen shows WEIGHT / DRUG MASS fields (body-weight mode)."),
 ("ICU Medical", "Medfusion 4000", "Syringe",
  "ICU page: 'Default into the drug library', 'Quick Library feature'.",
  "ICU page: 'Wireless connectivity with the PharmGuard Server enables updates to the drug library ... EMR integration'. Product name: 'Medfusion 4000 Wireless Syringe Infusion Pump'.",
  "Connected", "Third-party (India distributor)", IM_MF, "Features from ICU global page: " + ICU_MF, "Connected (peach)", "Confirmed",
  "IndiaMart listing verified by project team: seller Apex Medical India (Mark.Inc), Lucknow, Rs 1,80,000; product photo shows the 'PharmGuard Medication Safety' start screen (drug library). ICU Medical has no India product website."),
 ("ICU Medical", "Graseby 1200", "Volumetric",
  "None in India brochure (Smiths Medical, ref IN193908GB, lists Indian IV-set brands Romsons, Polymed).",
  "India brochure: 'Standard RS232 interface' only (not counted).",
  "Standard", "Third-party (India distributor)", MB_G1200,
  "IndiaMart: " + IM_G1200 + " | India brochure (Hospimax): " + G1200_BRO, "Standard (peach)", "Confirmed",
  "Medikabazaar (Rs 48,825) is primary - its photo shows the Graseby 1200 volumetric pump. IndiaMart listing (team-verified: Horizon Medical Technologies, New Delhi, Rs 55,000) is titled Graseby 1200 but its photo shows a Graseby 2000 SYRINGE pump - use as supporting only. Some non-Indian distributors list a newer 'Graseby 1200' with a 2,000-drug library - not seen on Indian sites."),
]

wb = Workbook()

def hdr(ws, row, values, fill=NAVY):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = fill; cell.alignment = CENTER; cell.border = BOX

def legend(ws, row, col):
    items = [(BLUE, "Listed on official India website"), (PEACH, "Listed on Indian distributor (third-party) website only"),
             (GREY, "Not disclosed on Indian websites")]
    for i, (fill, txt) in enumerate(items):
        ws.cell(row=row + i, column=col).fill = fill
        ws.cell(row=row + i, column=col).border = BOX
        ws.cell(row=row + i, column=col + 1, value=txt).font = Font(name=F, size=9)

# ---------------------------------------------------------------- Sheet 1: slide matrix
ws = wb.active; ws.title = "Slide Matrix"
ws["A1"] = "Task 2.2 | Infusion Pumps Market Assessment in India | Key Players (2/8) - corrected classification (Terumo, Mindray, ICU Medical)"
ws["A1"].font = Font(name=F, bold=True, size=13, color="0B1F6B")
ws["A2"] = "Criteria: Standard = no drug library, no connectivity | Smart = drug library, no connectivity | Smart** = drug library, connectivity optional (add-on module / WiFi variant) | Connected = drug library + integrated connectivity"
ws["A2"].font = Font(name=F, italic=True, size=9)
legend(ws, 3, 11)
hdr(ws, 4, ["Company", "", "", "Syringe Pump", "", "", "Volumetric Pump", "", "", "Changes vs sample"])
ws.merge_cells("A4:C4"); ws.merge_cells("D4:F4"); ws.merge_cells("G4:I4")
hdr(ws, 5, ["Name", "HQ", "Revenue", "Standard", "Smart", "Connected", "Standard", "Smart", "Connected", ""], fill=MID)
ws.merge_cells("J4:J5")

COLS = {("Syringe", "Standard"): 4, ("Syringe", "Smart"): 5, ("Syringe", "Smart**"): 5, ("Syringe", "Connected"): 6,
        ("Volumetric", "Standard"): 7, ("Volumetric", "Smart"): 8, ("Volumetric", "Smart**"): 8, ("Volumetric", "Connected"): 9}
companies = [
 ("Terumo", "Japan", "JPY 1,131.9B FY25 revenue (Apr-25 to Mar-26) (~USD 7.15B*)\nof which Medical Care Solutions (incl. infusion & syringe pumps): JPY 216.1B\nSource: Terumo FY2025 Financial Results",
  "- TE-SS730 & TE-LM730: Smart -> Standard (brochure: 'Standard pumps without IT functions'; library mode is 830-only)\n- TE-SS830 & TE-LM830 Connected confirmed (built-in wireless LAN)\n- 730 models appear only in the India-hosted brochure: confirm they are still sold"),
 ("Mindray", "China", "CNY 33.28B FY25 revenue (~USD 4.96B*)\nof which Life Information & Support (incl. infusion pumps): CNY 9.84B\nSource: Mindray 2025 Annual Report (SZSE filing)",
  "- 5 Series: Smart -> Smart** (wireless module optional; DS5 dock)\n- e Series (eSP/eVP): Smart -> Connected (DERS + wired/wireless CMS monitoring)\n- 1/3 Series Standard and u Series Connected confirmed\n- 3 Series has a 200-drug name list without dose limits: kept Standard (decision needed)"),
 ("ICU Medical", "United States", "USD 2.23B FY25 revenue\nof which Infusion Systems (global): USD 684.2M\nSource: ICU Medical Q4/FY2025 results release",
  "- No change in classification (Graseby 2000/2100 & 1200 Standard; Medfusion 4000 Connected; all distributor-only, peach)\n- Flag: ICU's FY2025 10-K does not mention Graseby, and graseby.com is now run by Lianying Graseby Medical (China): confirm who supplies Graseby in India"),
]
r = 6
for name, hq, rev, changes in companies:
    buckets = {}
    for p in P:
        if p[0] != name or p[6] == "Not offered in India" or p[5].startswith("Excluded"): continue
        col = COLS[(p[2], p[5])]
        label = p[1] + ("**" if p[5] == "Smart**" else "") + (" (pending India confirmation)" if p[6] == "Not verified in India" else "")
        buckets.setdefault(col, []).append((label, SRC_FILL[p[6]]))
    n = max(len(v) for v in buckets.values())
    for c in range(1, 11):
        for i in range(n):
            ws.cell(row=r + i, column=c).border = BOX
    for c, v in ((1, name), (2, hq), (3, rev), (10, changes)):
        ws.cell(row=r, column=c, value=v).font = Font(name=F, size=9, bold=(c == 1))
        ws.cell(row=r, column=c).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=c, end_row=r + n - 1, end_column=c)
    for col in range(4, 10):
        items = buckets.get(col, [])
        for i in range(n):
            cell = ws.cell(row=r + i, column=col)
            if i < len(items):
                cell.value, cell.fill = items[i]
            else:
                cell.fill = GREY if not items else PatternFill()
            cell.font = Font(name=F, size=9); cell.alignment = CENTER
        if not items:
            ws.merge_cells(start_row=r, start_column=col, end_row=r + n - 1, end_column=col)
            ws.cell(row=r, column=col).fill = GREY
    for i in range(n): ws.row_dimensions[r + i].height = 32
    r += n
ws.cell(row=r + 1, column=1, value="** Smart** = drug library + connectivity optional (optional wireless module / docking station).  "
        "* USD converted at ECB reference rates of 5-Oct-2026 (latest available): USD/JPY 158.23, USD/CNY 6.705 (cross rates via EUR).").font = Font(name=F, size=8, italic=True)
ws.cell(row=r + 2, column=1, value="Specialty anaesthesia (TCI) pumps are excluded from scope. Product-level evidence and URLs: see 'Product Evidence' tab.").font = Font(name=F, size=8, italic=True)
for c, w in zip("ABCDEFGHIJKL", [14, 12, 34, 24, 24, 24, 22, 22, 24, 60, 4, 48]):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "A6"

# ---------------------------------------------------------------- Sheet 2: product evidence
ws2 = wb.create_sheet("Product Evidence")
heads = ["Company", "Product", "Pump type", "Drug library - evidence (quoted)", "Connectivity - evidence (quoted)",
         "Classification", "India listing source", "Primary URL (India, client-viewable)", "Supporting URL(s) / brochure page",
         "Sample slide said", "Change vs sample", "Notes", "Evidence screenshot (evidence/ folder)"]
hdr(ws2, 1, heads)
ter_shots = "Terumo_IN_brochure_p2_Smart_vs_Standard.jpg; Terumo_IN_brochure_p4_IT_solution_specs.jpg; Terumo_IN_Terufusion_Advanced_Infusion.jpg"
shots = {"Terumo": ter_shots,
         "BeneFusion 1": "Mindray_IN_BeneFusion_1_Series.jpg",
         "BeneFusion 3": "Mindray_IN_BeneFusion_3_Series.jpg; Mindray_SP3_datasheet_IndiaMart.jpg",
         "BeneFusion 5": "Mindray_IN_BeneFusion_5_Series.jpg",
         "BeneFusion e": "Mindray_IN_BeneFusion_e_Series.jpg",
         "BeneFusion u": "Mindray_IN_BeneFusion_u_Series.jpg; Mindray_uSeries_brochure_IndiaMart_p2.jpg",
         "Graseby 2000": "IndiaMart_Graseby_2000_team_check.jpg; Honmed_Graseby_2100.jpg", "Graseby 2100": "Honmed_Graseby_2100.jpg; IndiaMart_Graseby_2100_team_check.jpg",
         "Medfusion 4000": "IndiaMart_Medfusion_4000_team_check.jpg; ICU_Global_Medfusion_4000.jpg",
         "Graseby 1200": "Medikabazaar_Graseby_1200.jpg; IndiaMart_Graseby_1200_team_check.jpg; Graseby_1200_India_brochure_Hospimax_p2.jpg; ..._p3.jpg"}
def shot(p):
    for k, v in shots.items():
        if p[0] == k or p[1].startswith(k): return v
    return ""
for i, p in enumerate(P, 2):
    row = list(p) + [shot(p)]
    for c, v in enumerate(row, 1):
        cell = ws2.cell(row=i, column=c, value=v)
        cell.font = Font(name=F, size=9, bold=(c == 6)); cell.alignment = WRAP; cell.border = BOX
    fill = SRC_FILL[p[6]]
    for c in (2, 7, 8): ws2.cell(row=i, column=c).fill = fill
    if p[7].startswith("http"): ws2.cell(row=i, column=8).hyperlink = p[7]; ws2.cell(row=i, column=8).font = Font(name=F, size=9, color="0563C1", underline="single")
for c, w in enumerate([13, 26, 11, 44, 50, 13, 18, 44, 55, 16, 32, 44, 32], 1):
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.freeze_panes = "C2"; ws2.auto_filter.ref = f"A1:M{len(P)+1}"

# ---------------------------------------------------------------- Sheet 3: revenue
ws3 = wb.create_sheet("Revenue")
ws3["A1"] = "Revenue cross-check (FY2025)"; ws3["A1"].font = Font(name=F, bold=True, size=12)
fx = [("EUR/USD (input)", 1.1204), ("EUR/JPY (input)", 177.28), ("EUR/CNY (input)", 7.5118)]
for i, (lbl, v) in enumerate(fx, 2):
    ws3.cell(row=i, column=1, value=lbl).font = Font(name=F, size=9)
    c = ws3.cell(row=i, column=2, value=v); c.font = Font(name=F, color="0000FF"); c.fill = INPUT
ws3["C2"] = "ECB euro reference rates, latest available (5-Oct-2026): https://data.ecb.europa.eu/data/datasets/EXR  - USD crosses derived below"
ws3["C3"] = "=\"Implied USD/JPY: \"&TEXT(B3/B2,\"0.00\")"
ws3["C4"] = "=\"Implied USD/CNY: \"&TEXT(B4/B2,\"0.000\")"
for c in ("C2", "C3", "C4"): ws3[c].font = Font(name=F, size=9, italic=True)
FXCELL = {"USD": "$B$2", "JPY": "$B$3", "CNY": "$B$4"}
hdr(ws3, 6, ["Company", "Metric", "Period", "Currency", "Value (local, mn)", "Value (USD bn) - formula",
             "Sample slide", "Check", "Source URL", "Page / location", "Comment"])
TER_FY = "https://www.terumo.com/system/files/document/2026-05/FinancialResults_26Q4_E.pdf"
MR_AR = "http://static.cninfo.com.cn/finalpage/2026-03-31/1225059012.PDF"
MR_AR_SUM = "http://static.cninfo.com.cn/finalpage/2026-03-31/1225059011.PDF"
ICU_PR = "https://ir.icumed.com/news-releases/news-release-details/icu-medical-announces-fourth-quarter-2025-results-and-provides"
rev = [
 ("Terumo", "Group revenue", "FY2025 (Apr-25 to Mar-26)", "JPY", 1131877, "~$7.2B (FY25)", "Matches (USD 7.15B at latest ECB cross rate)",
  TER_FY, "PDF p.1 (summary) and p.6: 'Revenue totaled JPY 1,131.9 billion, an increase of 9.2%'", "IFRS consolidated."),
 ("Terumo", "Medical Care Solutions Company revenue (incl. infusion & syringe pumps)", "FY2025", "JPY", 216138, "-", "Closest infusion proxy",
  TER_FY, "PDF p.7 (segment table); p.19 lists 'Syringes, Infusion pumps, Syringe pumps ...' under Hospital Care Solutions",
  "Also includes syringes, IV solutions, PD fluid, pharmaceutical solutions. Pump-only revenue not disclosed."),
 ("Mindray", "Group operating revenue", "FY2025 (Dec-25)", "CNY", 33282.1594, "~$4.9B (FY25)", "Matches (USD 4.96B at latest ECB cross rate)",
  MR_AR_SUM, "Annual report summary PDF p.3: '营业收入 3,328,215.94 万元，较上年同期下降 9.38%' (CNY 33,282mn, -9.4%)",
  "Official SZSE filing via CNINFO (Chinese). Full annual report: " + MR_AR + "  Sample's source (BigGo) is a news site."),
 ("Mindray", "Life Information & Support revenue (incl. infusion pumps)", "FY2025", "CNY", 9836.7237, "-", "Closest infusion proxy",
  MR_AR_SUM, "Annual report summary PDF p.9: '生命信息与支持业务实现营业收入 983,672.37 万元，同比下降 19.80%'; product list incl. '输注泵' (infusion pumps)",
  "Also includes monitors, anaesthesia, ventilators, defibrillators, OR lights/tables. Pump-only revenue not disclosed."),
 ("ICU Medical", "Total revenues", "FY2025 (Dec-25)", "USD", 2231.262, "~$2.2B (FY 2025)", "Matches",
  ICU_PR, "Results table: 'TOTAL REVENUES ... 2,231,262' (USD thousands, twelve months)", "Also in FY2025 10-K (SEC)."),
 ("ICU Medical", "Infusion Systems revenue (global)", "FY2025", "USD", 684.208, "~$684.2M (Global Infusion Pumps Revenue)", "Matches; label as 'Infusion Systems' (incl. LVP & syringe pumps, dedicated sets, software)",
  ICU_PR, "Revenue by product line table: 'Infusion Systems ... 684.2'", "Not pump-hardware-only: includes dedicated sets and software."),
]
for i, rw in enumerate(rev, 7):
    fxc = FXCELL[rw[3]]
    vals = list(rw[:5]) + [f"=IF(E{i}=\"\",\"n/a\",E{i}/({fxc}/$B$2)/1000)"] + list(rw[5:])
    for c, v in enumerate(vals, 1):
        cell = ws3.cell(row=i, column=c, value=v); cell.font = Font(name=F, size=9); cell.alignment = WRAP; cell.border = BOX
    ws3.cell(row=i, column=5).font = Font(name=F, size=9, color="0000FF"); ws3.cell(row=i, column=5).number_format = "#,##0.0"
    ws3.cell(row=i, column=6).number_format = "0.00"
    ws3.cell(row=i, column=9).hyperlink = rw[7]
for c, w in enumerate([13, 34, 18, 9, 14, 14, 22, 30, 50, 50, 50], 1):
    ws3.column_dimensions[get_column_letter(c)].width = w

# ---------------------------------------------------------------- Sheet 4: clean URL map
ws4 = wb.create_sheet("Clean URL Map")
hdr(ws4, 1, ["Company", "Product / purpose", "URL (client-facing)", "Source type", "Status", "Reason / note"])
urls = [
 ("Terumo", "Terufusion Advanced Infusion Systems (India product page)", TER_PAGE, "Official India website", "KEEP", "Names Smart Syringe Pump, Smart Infusion Pump, PMS software; 'integrated drug library'"),
 ("Terumo", "Smart Infusion System brochure (p.2 Smart vs Standard; p.4 specs)", TER_BRO, "Brochure on official India website", "KEEP", "Key evidence: 730 = 'Standard pumps without IT functions'; wireless LAN & library mode = 830 only"),
 ("Terumo", "Launch press release, Aug-2025 (India)", TER_PR, "Official India website", "KEEP", "Confirms India launch of the connected system"),
 ("Terumo", "FY2025 Financial Results (revenue)", TER_FY, "Company financial results", "KEEP", "p.1/p.6 group revenue; p.7 segment; p.19 segment products"),
 ("Mindray", "India - Infusion System range (lists 1, 3, 5, e, u Series)", MR, "Official India website", "KEEP", "Master list"),
 ("Mindray", "BeneFusion 1 Series", MR1, "Official India website", "KEEP", "Standard"),
 ("Mindray", "BeneFusion 3 Series", MR3, "Official India website", "ADDED", "Was missing from source list; Standard"),
 ("Mindray", "BeneFusion 5 Series", MR5, "Official India website", "ADDED", "DERS drug library + DS5/CS5 connectivity (Smart**)"),
 ("Mindray", "BeneFusion e Series", MRE, "Official India website", "KEEP", "SafeDose DERS + Integrated Central Monitoring"),
 ("Mindray", "BeneFusion u Series", MRU, "Official India website", "KEEP", "HL7 connectivity"),
 ("Mindray", "u Series brochure (IndiaMart-hosted) p.2", MR_U_BRO, "Third-party (India distributor)", "ADDED (peach)", "'Departmental drug library'"),
 ("Mindray", "SP3 data sheet (IndiaMart-hosted)", MR_SP3_DS, "Third-party (India distributor)", "ADDED (peach)", "'Drug library up to 200 drugs' - name list, no dose limits"),
 ("Mindray", "VP5 operator manual s.7.4", MR_VP5_OM, "Company manual (via distributor, Syria)", "ADDED (backup)", "'Wireless Networking (Optional)'"),
 ("Mindray", "eSP (PCA) data sheet", MR_ESP_DS, "Company data sheet (via distributor, S. Africa)", "ADDED (backup)", "'Communication: Wired/wireless'; drug library 5000 drugs"),
 ("Mindray", "2025 Annual Report summary (revenue, p.3 & p.9)", MR_AR_SUM, "Company annual report (SZSE / CNINFO)", "ADDED", "Official source; replaces news article"),
 ("Mindray", "BigGo Finance article (revenue)", "https://finance.biggo.com/news/Vmaddp0BOIb5XxavmPY-", "News aggregator", "DROP", "Unofficial; its segment table has unit errors (e.g., '122.41' for CNY 12.24B). Use annual report"),
 ("ICU Medical", "Medfusion 4000 (IndiaMart listing)", IM_MF, "Third-party (India distributor)", "KEEP (peach) - team-verified", "Apex Medical India, Lucknow, Rs 1,80,000"),
 ("ICU Medical", "Medfusion 4000 Wireless (ICU global page)", ICU_MF, "Global website", "KEEP (backup)", "Feature proof only: drug library + wireless PharmGuard"),
 ("ICU Medical", "Graseby 2000 (IndiaMart listing)", IM_G2000, "Third-party (India distributor)", "ADDED (peach) - team-verified", "Nature's Global Service, New Delhi, Rs 30,000. Sample had no URL for Graseby 2000/2100"),
 ("ICU Medical", "Graseby 2100 (Honmed, India)", HON_G2100, "Third-party (India distributor)", "ADDED (peach)", "Graseby 2000 range description; no drug library"),
 ("ICU Medical", "Graseby 2100 (IndiaMart listing)", IM_G2100, "Third-party (India distributor)", "ADDED (peach) - team-verified", "Apex Medical India, Lucknow, Rs 30,000"),
 ("ICU Medical", "Graseby 1200 (IndiaMart listing)", IM_G1200, "Third-party (India distributor)", "SUPPORTING ONLY (peach) - team-verified", "Horizon Medical Technologies, New Delhi, Rs 55,000. CAUTION: listing photo shows a Graseby 2000 syringe pump - show Medikabazaar to clients instead"),
 ("ICU Medical", "Graseby 1200 (Medikabazaar)", MB_G1200, "Third-party (India distributor)", "KEEP (peach) - PRIMARY", "Tracking parameters removed; photo shows the 1200 volumetric pump; Rs 48,825"),
 ("ICU Medical", "Graseby 1200 India brochure (Hospimax)", G1200_BRO, "Third-party (India distributor)", "ADDED (peach)", "India brochure: RS232 only, no drug library"),
 ("ICU Medical", "Q4/FY2025 results release (revenue)", ICU_PR, "Company results release", "KEEP", "Total USD 2,231.3mn; Infusion Systems USD 684.2mn"),
 ("ICU Medical", "IR static file (FY2025)", "https://ir.icumed.com/static-files/d27738f3-0a46-4a32-8d61-0efc3a3852bc", "Company IR document", "DROP", "130-page document; the results release above is enough for the two figures"),
]
for i, u in enumerate(urls, 2):
    for c, v in enumerate(u, 1):
        cell = ws4.cell(row=i, column=c, value=v); cell.font = Font(name=F, size=9); cell.alignment = WRAP; cell.border = BOX
    if u[2].startswith("http"):
        ws4.cell(row=i, column=3).hyperlink = u[2]; ws4.cell(row=i, column=3).font = Font(name=F, size=9, color="0563C1", underline="single")
    if "Third-party (India" in u[3]: ws4.cell(row=i, column=4).fill = PEACH
    elif "official India" in u[3] or "Official India" in u[3]: ws4.cell(row=i, column=4).fill = BLUE
    if u[4] == "DROP": ws4.cell(row=i, column=5).fill = RED
for c, w in enumerate([14, 44, 70, 30, 24, 60], 1):
    ws4.column_dimensions[get_column_letter(c)].width = w
ws4.freeze_panes = "A2"; ws4.auto_filter.ref = f"A1:F{len(urls)+1}"

# ---------------------------------------------------------------- Sheet 5: method & caveats
ws5 = wb.create_sheet("Method & Caveats")
notes = [
 "Scope: India market, Terumo, Mindray, ICU Medical (slide 2/8); general-purpose syringe and volumetric pumps. Specialty anaesthesia/TCI pumps are excluded; TCI modes of general pumps (e.g., SP5 TCI) are not counted separately. Checked October 2026.",
 "Source priority: (1) company India website, including brochures hosted on it; (2) Indian distributor websites incl. IndiaMart and IndiaMart-hosted brochures (imimg.com), marked peach; (3) global company pages/data sheets, used only as backup for features and never as proof of India availability.",
 "Classification rules (same as slide 1): optional wireless module / docking station / separate WiFi variant = Smart**. RS232 / nurse-call ports are NOT counted as connectivity. Body-weight dose calculation (Graseby 2100) is NOT a drug library.",
 "Judgement call (Mindray 3 Series): data sheets show a drug-NAME library (up to 200 drugs) with no dose limits. Kept Standard, as in the sample; it moves to Smart if any drug library counts.",
 "Terumo TE-SS730 / TE-LM730 appear only in the India-hosted brochure (2018 edition); the India web page text mentions only the Smart (830) pumps.",
 "Graseby ownership flag: ICU Medical's FY2025 10-K describes Medfusion 4000 but does not mention Graseby, and graseby.com is now operated by Lianying Graseby Medical (Zhejiang, China). Indian listings still say 'Smiths Medical Graseby'. Confirm Graseby's current owner / India supplier before keeping it under ICU Medical.",
 "IndiaMart blocks automated access from this environment (HTTP 429); all four IndiaMart listings (Medfusion 4000, Graseby 2000, 2100, 1200) were verified manually by the project team (screenshots in evidence/). smiths-medical.com (incl. /en-in) was unreachable.",
 "Revenue: official sources only on the slide (Terumo FY results, Mindray SZSE annual report, ICU results release). USD uses ECB reference rates of 5-Oct-2026 (latest), crossed via EUR: USD/JPY 158.23, USD/CNY 6.705 - editable inputs on the Revenue tab. No company reports pump-only revenue; segment figures are the closest proxies.",
]
ws5["A1"] = "Method, assumptions and caveats"; ws5["A1"].font = Font(name=F, bold=True, size=12)
for i, n in enumerate(notes, 3):
    ws5.cell(row=i, column=1, value=f"{i-2}. {n}").font = Font(name=F, size=10)
    ws5.cell(row=i, column=1).alignment = WRAP
    ws5.row_dimensions[i].height = 42
ws5.column_dimensions["A"].width = 150

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
